from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
import xlrd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def convert_xls_to_xlsx(source_path: Path, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    source_workbook = xlrd.open_workbook(source_path)
    target_workbook = openpyxl.Workbook()

    for sheet_index, source_sheet in enumerate(source_workbook.sheets()):
        if sheet_index == 0:
            target_sheet = target_workbook.active
            target_sheet.title = safe_sheet_title(source_sheet.name)
        else:
            target_sheet = target_workbook.create_sheet(safe_sheet_title(source_sheet.name))

        for row_index in range(source_sheet.nrows):
            for column_index in range(source_sheet.ncols):
                cell = source_sheet.cell(row_index, column_index)
                value = convert_xls_cell_value(cell, source_workbook.datemode)
                if value is not None:
                    target_sheet.cell(row=row_index + 1, column=column_index + 1, value=value)

    target_workbook.save(output_path)
    return output_path


def convert_xls_cell_value(cell: xlrd.sheet.Cell, datemode: int) -> object | None:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, datemode)
        if isinstance(value, datetime):
            return value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    if isinstance(cell.value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", cell.value)
    return cell.value


def safe_sheet_title(title: str) -> str:
    invalid_characters = set("[]:*?/\\")
    cleaned = "".join("_" if character in invalid_characters else character for character in title)
    cleaned = cleaned.strip() or "Sheet"
    return cleaned[:31]
